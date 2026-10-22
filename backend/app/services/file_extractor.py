import os
from typing import Any, Tuple
from pypdf import PdfReader
from docx import Document


class FileExtractionError(Exception):
    pass


def extract_text_from_pdf(file_path: Any) -> str:
    """Extracts text content from a PDF file using pypdf. Preserves layout, tables, and detects scanned/unreadable PDFs."""
    try:
        if isinstance(file_path, bytes):
            import io
            file_path = io.BytesIO(file_path)
        reader = PdfReader(file_path)
        if len(reader.pages) == 0:
            raise FileExtractionError("The uploaded PDF has 0 pages and contains no content.")

        extracted_text = []
        for index, page in enumerate(reader.pages):
            text = ""
            # First attempt layout-aware extraction to preserve columns and tables
            try:
                text = page.extract_text(extraction_mode="layout")
            except Exception:
                pass
            if not text or not text.strip():
                try:
                    text = page.extract_text()
                except Exception:
                    pass
            if text and text.strip():
                extracted_text.append(text.strip())

        combined = "\n\n".join(extracted_text).strip()
        # If pages exist but no text or very few characters could be extracted, it's a scanned image
        alphanumeric_chars = [c for c in combined if c.isalnum()]
        if len(alphanumeric_chars) < 30:
            raise FileExtractionError(
                "The uploaded PDF appears to be a scanned image or contains no selectable text. "
                "Optical character recognition (OCR) or a text-based PDF/DOCX file is required for automated profile extraction."
            )

        return combined
    except FileExtractionError:
        raise
    except Exception as e:
        raise FileExtractionError(f"Could not read PDF file: {str(e)}")


def extract_text_from_docx(file_path: str) -> str:
    """Extracts text content from a DOCX file including paragraphs, bullet lists, and tables."""
    try:
        doc = Document(file_path)
        extracted_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                extracted_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                seen_cells = set()
                row_text = []
                for cell in row.cells:
                    ctext = cell.text.strip()
                    if ctext and ctext not in seen_cells:
                        seen_cells.add(ctext)
                        row_text.append(ctext)
                if row_text:
                    extracted_text.append(" | ".join(row_text))

        combined = "\n".join(extracted_text).strip()
        alphanumeric_chars = [c for c in combined if c.isalnum()]
        if len(alphanumeric_chars) < 20:
            raise FileExtractionError(
                "The uploaded Word document contains no readable text. "
                "Please verify the file content or upload a different resume."
            )

        return combined
    except FileExtractionError:
        raise
    except Exception as e:
        raise FileExtractionError(f"Could not read Word (.docx) document: {str(e)}")


def extract_text_from_file(file_path: str, filename: str) -> Tuple[str, str]:
    """
    Extracts text from PDF or DOCX file.
    Returns (extracted_text, file_extension).
    Raises FileExtractionError with clear messages for scanned or unsupported files.
    """
    if not os.path.exists(file_path):
        raise FileExtractionError("File not found on storage.")
    
    ext = os.path.splitext(filename.lower())[1]
    
    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        text = extract_text_from_docx(file_path)
    else:
        raise FileExtractionError(
            f"Unsupported file format '{ext}'. Only PDF (.pdf) and Microsoft Word (.docx) files are supported."
        )
    
    return text, ext
