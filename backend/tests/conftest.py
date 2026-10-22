import io
import os
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from docx import Document
from pypdf import PdfWriter

from app.core.database import Base, get_db
from app.core.security import create_access_token, hash_password
from app.main import app
from app.models.user import User

# In-memory SQLite database for isolated test runs
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


@pytest.fixture(scope="function")
def db_session():
    Base.metadata.create_all(bind=engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture(scope="function")
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def test_user(db_session):
    user = User(
        email="testuser@example.com",
        hashed_password=hash_password("securepassword123"),
        full_name="Alex Morgan",
        font_size="normal",
        display_mode="simple",
        is_active=True
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


@pytest.fixture
def auth_headers(test_user):
    token = create_access_token(subject=test_user.id)
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def sample_docx_bytes():
    """Generates a valid minimal DOCX in memory with sample resume text."""
    doc = Document()
    doc.add_heading("Alex Morgan", 0)
    doc.add_paragraph("alex.morgan@example.com | (555) 345-6789 | San Francisco, CA")
    doc.add_paragraph("https://linkedin.com/in/alexmorgan | https://github.com/alexmorgan")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("Experienced Full Stack Software Engineer passionate about accessible, user-friendly web applications.")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, React, TypeScript, FastAPI, Docker, SQL, Git")
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Senior Software Engineer at TechCorp (2021 - Present)")
    doc.add_paragraph("Led architecture of modern web applications using React and Python.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science, University of California (2020)")
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream.getvalue()


@pytest.fixture
def sample_pdf_bytes():
    """Generates a valid minimal PDF in memory."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    stream = io.BytesIO()
    writer.write(stream)
    stream.seek(0)
    return stream.getvalue()
